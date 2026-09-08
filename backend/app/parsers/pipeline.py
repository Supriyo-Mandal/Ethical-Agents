from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

from .detector import detect_format
from .models import block, document_model, finalize


def _ocr(raw: bytes, document: dict[str, Any]) -> str:
    try:
        from PIL import Image
        import pytesseract
        text = pytesseract.image_to_string(Image.open(io.BytesIO(raw))).strip()
        document["metadata"]["ocr_used"] = True
        return text
    except ImportError:
        document["warnings"].append("OCR dependencies are unavailable")
    except Exception as exc:
        document["warnings"].append(f"OCR failed: {exc.__class__.__name__}")
    return ""


def _text_page(number: int, blocks: list[dict[str, Any]]) -> dict[str, Any]:
    return {"page_number": number, "blocks": blocks}


def _parse_pdf(raw: bytes, document: dict[str, Any]) -> None:
    try:
        import fitz
    except ImportError as exc:
        raise ValueError("PDF support requires pymupdf") from exc
    pdf = fitz.open(stream=raw, filetype="pdf")
    for page_number, page in enumerate(pdf, 1):
        blocks: list[dict[str, Any]] = []
        for index, item in enumerate(page.get_text("blocks")):
            text = item[4].strip()
            if text:
                blocks.append(block(f"p{page_number}-text{index}", "paragraph", text, bbox=item[:4], source="native"))
        if not blocks:
            image = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False).tobytes("png")
            text = _ocr(image, document)
            blocks.append(block(f"p{page_number}-ocr", "ocr", text, source="ocr"))
        for index, image in enumerate(page.get_images(full=True)):
            try:
                extracted = pdf.extract_image(image[0])["image"]
                ocr_text = _ocr(extracted, document)
                blocks.append(block(f"p{page_number}-image{index}", "figure", ocr_text, source="ocr", image=True))
            except Exception:
                document["warnings"].append(f"Could not extract image on page {page_number}")
        document["pages"].append(_text_page(page_number, blocks))


def _parse_docx(raw: bytes, document: dict[str, Any]) -> None:
    try:
        import docx
    except ImportError as exc:
        raise ValueError("DOCX support requires python-docx") from exc
    parsed = docx.Document(io.BytesIO(raw))
    blocks: list[dict[str, Any]] = []
    for index, paragraph in enumerate(parsed.paragraphs):
        if paragraph.text.strip():
            style_name = str(paragraph.style.name or "") if paragraph.style is not None else ""
            kind = "heading" if style_name.startswith("Heading") else "paragraph"
            blocks.append(block(f"p1-text{index}", kind, paragraph.text))
    for index, table in enumerate(parsed.tables):
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        blocks.append(block(f"p1-table{index}", "table", "\n".join(rows)))
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        for index, name in enumerate(n for n in archive.namelist() if n.startswith("word/media/")):
            text = _ocr(archive.read(name), document)
            blocks.append(block(f"p1-image{index}", "figure", text, source="ocr", image=True))
    document["pages"].append(_text_page(1, blocks))


def _parse_markdown(raw: bytes, document: dict[str, Any]) -> None:
    text = raw.decode("utf-8", errors="replace")
    blocks = []
    for index, line in enumerate(text.splitlines()):
        if line.strip():
            kind = "heading" if line.lstrip().startswith("#") else "paragraph"
            blocks.append(block(f"p1-md{index}", kind, line.strip()))
    document["pages"].append(_text_page(1, blocks))


def parse_upload(raw: bytes, filename: str, content_type: str = "application/octet-stream") -> dict[str, Any]:
    if not raw:
        raise ValueError("Uploaded file is empty")
    format_name = detect_format(filename, raw)
    document = document_model(filename, content_type, format_name)
    if format_name in {"txt", "md"}:
        _parse_markdown(raw, document) if format_name == "md" else document["pages"].append(_text_page(1, [block("p1-text0", "paragraph", raw.decode("utf-8", errors="replace"))]))
    elif format_name == "pdf":
        _parse_pdf(raw, document)
    elif format_name == "docx":
        _parse_docx(raw, document)
    else:
        document["pages"].append(_text_page(1, [block("p1-image0", "image", _ocr(raw, document), source="ocr", image=True)]))
    return finalize(document)
